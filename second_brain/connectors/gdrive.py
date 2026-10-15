"""Google Drive connector: files in the folders listed in sources.yaml (and their subfolders), never the whole drive.

First run (no cursor): walk the folders and index every supported file.
Later runs: the Drive Changes API since the saved page token; files trashed, deleted or moved out of the
folders become Deletions. Spreadsheets (.xlsx, Google Sheets, CSV) become one "| a | b |" line per row.
Folders in exclude_folder_ids (and everything below them) are skipped.
Files under a folder in private_folder_ids are tagged sensitivity: private (D-017).
Google Docs are exported as Markdown, PDFs and .docx are converted to text.
Scanned PDFs (no text layer) are skipped.

Login once with: uv run python -m second_brain login gdrive
"""

import io
import logging
from collections.abc import Iterator
from pathlib import Path

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from langchain_core.documents import Document

from second_brain.config import ROOT, sources
from second_brain.models import Deletion

log = logging.getLogger(__name__)
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
TOKEN_FILE = ROOT / "tokens" / "gdrive.token.json"
FOLDER = "application/vnd.google-apps.folder"
MIME = {
    "gdoc": "application/vnd.google-apps.document",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown",
    "txt": "text/plain",
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "gsheet": "application/vnd.google-apps.spreadsheet",
}
MAX_ROWS = 2000  # per sheet; bigger sheets are usually exports, not notes
RETRIES = 5  # Drive returns occasional 500s and 429s; the client backs off exponentially between tries
MAX_BYTES = 20 * 1024 * 1024  # skip anything bigger; usually scans or exports, not notes
FILE_FIELDS = "id, name, mimeType, parents, webViewLink, createdTime, modifiedTime, size, trashed"


def credentials(interactive: bool = False) -> Credentials:
    """Saved login, refreshed when needed. interactive=True runs the one-time browser login."""
    creds = Credentials.from_authorized_user_file(str(TOKEN_FILE), SCOPES) if TOKEN_FILE.exists() else None
    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())
    elif not (creds and creds.valid):
        if not interactive:
            raise SystemExit("No Google Drive login yet: run `uv run python -m second_brain login gdrive`")
        client_file = ROOT / sources()["gdrive"]["credentials_file"]
        flow = InstalledAppFlow.from_client_secrets_file(str(client_file), SCOPES)
        creds = flow.run_local_server(port=8085, open_browser=False)
    TOKEN_FILE.parent.mkdir(exist_ok=True)
    TOKEN_FILE.write_text(creds.to_json())
    TOKEN_FILE.chmod(0o600)
    return creds


class GDriveConnector:
    name = "gdrive"

    def __init__(self):
        cfg = sources()["gdrive"]
        self.folder_ids = set(cfg.get("folder_ids") or [])
        if not self.folder_ids:
            raise SystemExit("gdrive.folder_ids is empty in config/sources.yaml")
        self.exclude_ids = set(cfg.get("exclude_folder_ids") or [])
        self.private_ids = set(cfg.get("private_folder_ids") or [])
        self.mimes = {MIME[t]: t for t in cfg.get("file_types") or MIME}
        self.drive = build("drive", "v3", credentials=credentials(), cache_discovery=False)
        self._ancestors: dict[str, frozenset[str]] = {}
        self._next: str | None = None

    # --- Connector protocol ------------------------------------------------------------------
    def fetch_changes(self, cursor: str | None) -> Iterator[Document | Deletion]:
        if not cursor:
            # take the token first, so edits made during the walk show up in the next sync
            self._next = self.drive.changes().getStartPageToken().execute(num_retries=RETRIES)["startPageToken"]
            for f in self._walk():
                if doc := self._to_document(f):
                    yield doc
            return
        token = cursor
        while token:
            resp = self.drive.changes().list(pageToken=token, includeRemoved=True, spaces="drive",
                                             fields=f"nextPageToken, newStartPageToken, changes(fileId, removed, file({FILE_FIELDS}))").execute(num_retries=RETRIES)
            for ch in resp["changes"]:
                f = ch.get("file")
                if ch.get("removed") or not f or f.get("trashed") or not self._under_folders(f):
                    yield Deletion(source="gdrive", source_id=ch["fileId"])  # a no-op if it was never indexed
                elif f["mimeType"] in self.mimes and (doc := self._to_document(f)):
                    yield doc
            token = resp.get("nextPageToken")
            self._next = resp.get("newStartPageToken") or self._next

    def next_cursor(self) -> str:
        return self._next or ""

    # --- Listing -----------------------------------------------------------------------------
    def _list(self, q: str) -> Iterator[dict]:
        page = None
        while True:
            resp = self.drive.files().list(q=q, pageSize=200, pageToken=page,
                                           fields=f"nextPageToken, files({FILE_FIELDS})").execute(num_retries=RETRIES)
            yield from resp["files"]
            if not (page := resp.get("nextPageToken")):
                return

    def _walk(self) -> Iterator[dict]:
        todo, seen = list(self.folder_ids), set()
        while todo:
            folder = todo.pop()
            if folder in seen or folder in self.exclude_ids:
                continue
            seen.add(folder)
            for f in self._list(f"'{folder}' in parents and trashed = false"):
                if f["mimeType"] == FOLDER:
                    todo.append(f["id"])
                elif f["mimeType"] in self.mimes:
                    yield f

    def _ancestors_of(self, f: dict) -> frozenset[str]:
        """All folder IDs above a file (walks parents through the API, cached per folder)."""
        found: set[str] = set()
        for parent in f.get("parents") or []:
            if parent not in self._ancestors:
                p = self.drive.files().get(fileId=parent, fields="id, parents").execute(num_retries=RETRIES)
                self._ancestors[parent] = frozenset({parent}) | self._ancestors_of(p)
            found |= self._ancestors[parent]
        return frozenset(found)

    def _under_folders(self, f: dict) -> bool:
        above = self._ancestors_of(f)
        return bool(above & self.folder_ids) and not above & self.exclude_ids

    # --- File -> Document --------------------------------------------------------------------
    def _to_document(self, f: dict) -> Document | None:
        if int(f.get("size") or 0) > MAX_BYTES:
            log.warning("skipping %s: larger than %d MB", f["name"], MAX_BYTES // 2**20)
            return None
        try:
            text = self._text(f).strip()
        except Exception as e:  # one unreadable file must not stop the sync
            log.warning("skipping %s: %s", f["name"], e)
            return None
        if not text:
            log.info("skipping %s: no text (scanned PDF or empty)", f["name"])
            return None
        log.info("gdrive file: %s (%d chars)", f["name"], len(text))
        title = f["name"] if self.mimes[f["mimeType"]] in ("gdoc", "gsheet") else Path(f["name"]).stem
        if not text.startswith("# "):
            text = f"# {title}\n\n{text}"
        return Document(page_content=text, metadata={
            "source": "gdrive",
            "source_id": f["id"],
            "title": title,
            "url": f.get("webViewLink", ""),
            "created_at": f["createdTime"],
            "updated_at": f["modifiedTime"],
            "sensitivity": "private" if self._ancestors_of(f) & self.private_ids else "normal",
        })

    def _text(self, f: dict) -> str:
        kind = self.mimes[f["mimeType"]]
        if kind == "gdoc":
            return self.drive.files().export(fileId=f["id"], mimeType="text/markdown").execute(num_retries=RETRIES).decode()
        if kind == "gsheet":  # exported as .xlsx so every tab comes along (CSV export is the first tab only)
            data = self.drive.files().export(fileId=f["id"], mimeType=MIME["xlsx"]).execute(num_retries=RETRIES)
            return xlsx_text(data)
        data = self._download(f["id"])
        if kind in ("md", "txt"):
            return data.decode("utf-8", errors="replace")
        if kind == "csv":
            import csv
            rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="replace"))))[:MAX_ROWS]
            return "\n".join(table_line(r) for r in rows if any(c.strip() for c in r))
        if kind == "xlsx":
            return xlsx_text(data)
        if kind == "pdf":
            from pypdf import PdfReader
            return "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        from docx import Document as Docx
        lines = []
        for para in Docx(io.BytesIO(data)).paragraphs:
            style = para.style.name if para.style else ""
            level = int(style.split()[-1]) if style.startswith("Heading") and style.split()[-1].isdigit() else 0
            lines.append(("#" * min(level, 3) + " " if level else "") + para.text)
        return "\n\n".join(line for line in lines if line.strip())

    def _download(self, file_id: str) -> bytes:
        buf = io.BytesIO()
        dl = MediaIoBaseDownload(buf, self.drive.files().get_media(fileId=file_id))
        done = False
        while not done:
            _, done = dl.next_chunk(num_retries=RETRIES)
        return buf.getvalue()


def table_line(cells) -> str:
    return "| " + " | ".join(str(c).strip() for c in cells) + " |"


def xlsx_text(data: bytes) -> str:
    """Each sheet as a "## <sheet>" section with one "| a | b |" line per non-empty row (values, not formulas)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else v for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                lines.append(table_line(cells))
            if len(lines) >= MAX_ROWS:
                break
        if lines:
            parts.append(f"## {ws.title}\n\n" + "\n".join(lines))
    return "\n\n".join(parts)


def find_folders(name: str | None = None) -> list[dict]:
    """Folders matching a name (or the top-level ones), to pick folder_ids for sources.yaml."""
    drive = build("drive", "v3", credentials=credentials(), cache_discovery=False)
    q = f"mimeType = '{FOLDER}' and trashed = false and " + (
        f"name contains '{name.replace(chr(39), chr(92) + chr(39))}'" if name else "'root' in parents")
    return drive.files().list(q=q, pageSize=100, fields="files(id, name, parents, modifiedTime)").execute(num_retries=RETRIES)["files"]
